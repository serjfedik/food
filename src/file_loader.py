"""Достаём текст из PDF / DOCX / изображений. Если PDF — скан, падаем в OCR."""
from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Optional

from . import ocr


SUPPORTED_EXTS = (".pdf", ".docx", ".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp", ".txt")


@dataclass
class LoadResult:
    text: str
    source_type: str        # 'pdf-text' | 'pdf-ocr' | 'docx' | 'image-ocr' | 'txt'
    ocr_used: bool
    pages: int = 0
    warnings: list[str] = None  # type: ignore[assignment]

    def __post_init__(self) -> None:
        if self.warnings is None:
            self.warnings = []


def _read_pdf_text(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return ("\n\n".join(p for p in parts if p).strip(), len(reader.pages))


def _read_docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    chunks: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            chunks.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks).strip()


def load_file(filename: str, data: bytes, force_ocr: bool = False, langs: Optional[str] = None) -> LoadResult:
    """Универсальный лоадер. filename — для определения типа, data — байты."""
    name = filename.lower()
    langs = langs or ocr.DEFAULT_LANGS

    if name.endswith(".pdf"):
        text, pages = ("", 0)
        warnings: list[str] = []
        if not force_ocr:
            try:
                text, pages = _read_pdf_text(data)
            except Exception as exc:
                warnings.append(f"pypdf: {exc}")
        meaningful = len(text.replace(" ", "").replace("\n", "")) >= 40
        if force_ocr or not meaningful:
            try:
                ocr_text = ocr.pdf_bytes_to_text_via_ocr(data, langs=langs)
                return LoadResult(
                    text=ocr_text,
                    source_type="pdf-ocr",
                    ocr_used=True,
                    pages=pages,
                    warnings=warnings,
                )
            except ocr.OCRUnavailable as exc:
                warnings.append(str(exc))
                return LoadResult(text=text, source_type="pdf-text", ocr_used=False, pages=pages, warnings=warnings)
        return LoadResult(text=text, source_type="pdf-text", ocr_used=False, pages=pages, warnings=warnings)

    if name.endswith(".docx"):
        return LoadResult(text=_read_docx_text(data), source_type="docx", ocr_used=False)

    if name.endswith((".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp")):
        warnings = []
        try:
            text = ocr.image_bytes_to_text(data, langs=langs)
            return LoadResult(text=text, source_type="image-ocr", ocr_used=True, warnings=warnings)
        except ocr.OCRUnavailable as exc:
            warnings.append(str(exc))
            return LoadResult(text="", source_type="image-ocr", ocr_used=False, warnings=warnings)

    if name.endswith(".txt"):
        try:
            return LoadResult(text=data.decode("utf-8"), source_type="txt", ocr_used=False)
        except UnicodeDecodeError:
            return LoadResult(text=data.decode("cp1251", errors="replace"), source_type="txt", ocr_used=False)

    raise ValueError(f"Неподдерживаемый формат: {filename}")
